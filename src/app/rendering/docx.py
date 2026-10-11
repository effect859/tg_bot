"""DOCX отчёта по практике: DocumentContent → байты .docx (python-docx).

Рендерер чистый: без БД, сети, LLM, Telegram и FileStorage. ``render`` синхронный
и занимает CPU — из async-кода его вызывают через ``asyncio.to_thread``.

Оформление — ГОСТ 7.32-2017, все значения в ``GOST_7_32`` (``Layout``):
A4, поля, Times New Roman 14, интервал 1,5, абзацный отступ, выравнивание по
ширине, номер страницы внизу по центру. Параметры заданы стилями (Normal,
Heading 1, Heading 2), а не прямым форматированием абзацев; темы шрифтов и цвета
шаблона python-docx заменяются явными значениями.

Структура: титульный лист (если есть title_page) — страница 1 без номера; затем
«СОДЕРЖАНИЕ» с полем TOC (уровни заголовков 1–2), разделы — каждый с новой
страницы, и «СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ» (если есть references). Нумерация
«1.», «1.1.» — по положению основного раздела (heading_numbers); «Введение»,
«Заключение» и список источников — структурные элементы ГОСТ: заглавными, по
центру, без номера. Номер входит в текст заголовка, поэтому виден в оглавлении и в
любом редакторе. Номера страниц оглавления рендерер не считает: поле TOC заполняет
Word или LibreOffice при открытии или обновлении полей.

Титульный лист — поля TITLE_PAGE_FIELDS; пустое поле выводится линией для
заполнения от руки, неизвестное поле отклоняется до создания документа.

Строки оглавления оформлены явными стилями toc 1 и toc 2. Из пакета убраны
артефакты шаблона python-docx (миниатюра, свойства приложения, customXml,
stylesWithEffects, w14:docId).

Результат детерминирован: одинаковый документ — одинаковые байты (метаданные без
дат и автора, у записей ZIP фиксированное время).
"""

import io
import zipfile
from dataclasses import dataclass

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Cm, Mm, Pt, RGBColor
from docx.styles.style import ParagraphStyle
from docx.text.paragraph import Paragraph

from app.domain.enums import DocType
from app.domain.report_document import (
    DocumentErrorCode,
    InvalidDocumentError,
    heading_numbers,
    validate_document_content,
)
from app.domain.title_page import TITLE_PAGE_FIELDS
from app.generation.schemas import DocumentContent, Section

DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass(frozen=True)
class Layout:
    """Параметры оформления документа — единственное место, где они заданы."""

    page_width_mm: float
    page_height_mm: float
    margin_left_mm: float
    margin_right_mm: float
    margin_top_mm: float
    margin_bottom_mm: float
    font_name: str
    font_size_pt: float
    line_spacing: float
    first_line_indent_cm: float
    # Отступ после заголовка; перед подразделом — тот же отступ.
    heading_space_pt: float
    language: str
    toc_title: str
    # Уровни заголовков в оглавлении (ключ \o поля TOC).
    toc_levels: str
    # Левый отступ строки подраздела в оглавлении (стиль toc 2).
    toc_level_indent_cm: float

    @property
    def text_width_mm(self) -> float:
        """Ширина текста: страница без левого и правого полей."""
        return self.page_width_mm - self.margin_left_mm - self.margin_right_mm


GOST_7_32 = Layout(
    page_width_mm=210,
    page_height_mm=297,
    margin_left_mm=30,
    margin_right_mm=15,
    margin_top_mm=20,
    margin_bottom_mm=20,
    font_name="Times New Roman",
    font_size_pt=14,
    line_spacing=1.5,
    first_line_indent_cm=1.25,
    heading_space_pt=12,
    language="ru-RU",
    toc_title="СОДЕРЖАНИЕ",
    toc_levels="1-2",
    toc_level_indent_cm=1.25,
)

REFERENCES_TITLE = "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ"
# Линия вместо данных, которых нет: титульный лист заполняется от руки.
BLANK_LINE = "_" * 30
# Длинные вуз, место практики и тема — мельче, чтобы титульный лист уместился на одной
# странице при предельных значениях ТЗ (проверено в Word). Текст не сокращается.
TITLE_PAGE_LONG_VALUE_CHARS = 150
TITLE_PAGE_LONG_VALUE_FONT_PT = 12

# Время записей ZIP: самое раннее, допустимое форматом.
_ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)
_THEME_FONT_ATTRIBUTES = ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme")


class PracticeReportDocxRenderer:
    """DocumentRenderer для отчёта по практике."""

    file_extension = "docx"
    mime_type = DOCX_MIME_TYPE

    def __init__(self, layout: Layout = GOST_7_32) -> None:
        self._layout = layout

    def render(self, content: DocumentContent) -> bytes:
        """Raises InvalidDocumentError — до создания документа."""
        if content.doc_type is not DocType.PRACTICE_REPORT:
            raise InvalidDocumentError(DocumentErrorCode.UNSUPPORTED_DOC_TYPE, "doc_type")
        # Поля титульного листа знает этот рендерер: незнакомое не теряется молча.
        for index, key in enumerate(content.title_page):
            if key not in TITLE_PAGE_FIELDS:
                raise InvalidDocumentError(
                    DocumentErrorCode.TITLE_PAGE_UNSUPPORTED, f"title_page.keys[{index}]"
                )
        validate_document_content(content)

        has_title_page = bool(content.title_page)
        document = Document()
        self._setup_styles(document)
        self._add_toc_styles(document)
        self._setup_page(document, title_page=has_title_page)
        if has_title_page:
            self._add_title_page(document, content.title_page)
        self._add_table_of_contents(document, new_page=has_title_page)
        numbers = heading_numbers([section.title for section in content.sections])
        for number, section in zip(numbers, content.sections, strict=True):
            self._add_section(document, section, number, level=1)
        if content.references:
            self._add_references(document, content.references)
        _set_metadata(document, content.title)
        _strip_template_artifacts(document)

        buffer = io.BytesIO()
        document.save(buffer)
        return _normalize_zip(buffer.getvalue())

    # --- Оформление ------------------------------------------------------------------

    def _setup_styles(self, document: DocxDocument) -> None:
        layout = self._layout
        # Умолчания документа: шрифт и язык вместо тем шаблона.
        defaults = document.styles.element.find(qn("w:docDefaults"))
        if defaults is not None:
            run_defaults = defaults.find(f"{qn('w:rPrDefault')}/{qn('w:rPr')}")
            if run_defaults is not None:
                self._set_fonts(run_defaults)
                _set_language(run_defaults, layout.language)

        normal = document.styles["Normal"]
        assert isinstance(normal, ParagraphStyle)
        normal.font.name = layout.font_name
        normal.font.size = Pt(layout.font_size_pt)
        self._set_fonts(normal.element.get_or_add_rPr())
        paragraph = normal.paragraph_format
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.first_line_indent = Cm(layout.first_line_indent_cm)
        paragraph.line_spacing = layout.line_spacing
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)

        for level in (1, 2):
            style = document.styles[f"Heading {level}"]
            assert isinstance(style, ParagraphStyle)
            style.font.name = layout.font_name
            style.font.size = Pt(layout.font_size_pt)
            style.font.bold = True
            style.font.italic = False
            style.font.color.rgb = RGBColor(0, 0, 0)
            self._set_fonts(style.element.get_or_add_rPr())
            heading = style.paragraph_format
            heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
            heading.first_line_indent = Cm(layout.first_line_indent_cm)
            heading.line_spacing = layout.line_spacing
            heading.space_before = Pt(0 if level == 1 else layout.heading_space_pt)
            heading.space_after = Pt(layout.heading_space_pt)
            heading.keep_with_next = True
            # Раздел верхнего уровня — с новой страницы.
            heading.page_break_before = level == 1

    def _set_fonts(self, run_properties: BaseOxmlElement) -> None:
        fonts = run_properties.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            run_properties.insert(0, fonts)
        for attribute in _THEME_FONT_ATTRIBUTES:
            fonts.attrib.pop(qn(attribute), None)
        for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attribute), self._layout.font_name)

    def _add_toc_styles(self, document: DocxDocument) -> None:
        """Встроенные стили строк оглавления toc 1 и toc 2.

        Без них Word берёт латентные стили на основе Normal, и строки оглавления
        получают абзацный отступ и выравнивание по ширине. Шрифт и интервал
        наследуются от Normal; номер страницы — у правого края текста, с точками.
        """
        layout = self._layout
        for level, left_indent_cm in ((1, 0.0), (2, layout.toc_level_indent_cm)):
            style = document.styles.add_style(f"toc {level}", WD_STYLE_TYPE.PARAGRAPH, builtin=True)
            assert isinstance(style, ParagraphStyle)
            style.base_style = document.styles["Normal"]
            paragraph = style.paragraph_format
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.first_line_indent = Cm(0)
            paragraph.left_indent = Cm(left_indent_cm)
            paragraph.tab_stops.add_tab_stop(
                Mm(layout.text_width_mm), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
            )

    def _setup_page(self, document: DocxDocument, *, title_page: bool) -> None:
        layout = self._layout
        for section in document.sections:
            section.orientation = WD_ORIENT.PORTRAIT
            section.page_width = Mm(layout.page_width_mm)
            section.page_height = Mm(layout.page_height_mm)
            section.left_margin = Mm(layout.margin_left_mm)
            section.right_margin = Mm(layout.margin_right_mm)
            section.top_margin = Mm(layout.margin_top_mm)
            section.bottom_margin = Mm(layout.margin_bottom_mm)
            footer = section.footer.paragraphs[0]
            _centered(footer)
            _add_field(footer, "PAGE")
            if title_page:
                # Титульный лист — страница 1 без номера (w:titlePg и пустой колонтитул
                # первой страницы); со второй страницы номер есть.
                section.different_first_page_header_footer = True
                section.first_page_footer.is_linked_to_previous = False

    # --- Содержимое ------------------------------------------------------------------

    def _add_title_page(self, document: DocxDocument, fields: dict[str, str]) -> None:
        """Типовой титульный лист: известное — из ТЗ, остальное — линии для заполнения."""

        def value(key: str) -> str:
            return fields.get(key) or BLANK_LINE

        def long(key: str) -> bool:
            return len(fields.get(key) or "") > TITLE_PAGE_LONG_VALUE_CHARS

        def line(
            text: str,
            *,
            left: bool = False,
            bold: bool = False,
            before: int = 0,
            small: bool = False,
        ) -> None:
            paragraph = document.add_paragraph(text)
            paragraph_format = paragraph.paragraph_format
            paragraph_format.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.CENTER
            )
            paragraph_format.first_line_indent = Cm(0)
            paragraph_format.line_spacing = 1.0
            paragraph_format.space_before = Pt(before)
            if bold:
                paragraph.runs[0].bold = True
            if small:
                paragraph.runs[0].font.size = Pt(TITLE_PAGE_LONG_VALUE_FONT_PT)

        line(value("university"), small=long("university"))
        line(f"Факультет (институт): {value('faculty')}")
        line(f"Кафедра: {value('department')}")
        line("ОТЧЁТ", bold=True, before=36)
        line("о прохождении практики", bold=True)
        line(f"Вид практики: {value('practice_type')}", before=12)
        line(
            f"Место прохождения практики: {value('practice_place')}",
            small=long("practice_place"),
        )
        line(f"Тема: {value('topic')}", small=long("topic"))
        line(f"Сроки практики: {value('practice_period')}")
        line(f"Студент: {value('student_name')}", left=True, before=24)
        line(f"Группа: {value('student_group')}", left=True)
        line(f"Направление подготовки: {value('specialty')}", left=True)
        line(f"Руководитель практики от вуза: {value('supervisor')}", left=True)
        line(
            f"Руководитель практики от организации: {value('organization_supervisor')}",
            left=True,
        )
        line(f"{value('city')} {value('year')}", before=24)

    def _add_table_of_contents(self, document: DocxDocument, *, new_page: bool) -> None:
        title = document.add_paragraph(self._layout.toc_title)
        _centered(title)
        title.runs[0].bold = True
        title.paragraph_format.page_break_before = new_page
        field = document.add_paragraph()
        field.paragraph_format.first_line_indent = Cm(0)
        # \o — уровни заголовков, \h — ссылки, \z — без номеров страниц в веб-виде,
        # \u — уровни абзацев. Поле помечено «устаревшим»: Word обновит его при открытии.
        _add_field(field, f'TOC \\o "{self._layout.toc_levels}" \\h \\z \\u', dirty=True)
        _update_fields_on_open(document)

    def _add_section(
        self, document: DocxDocument, section: Section, number: str | None, *, level: int
    ) -> None:
        """Раздел с номером или — ``number`` None — структурный элемент без номера.

        Заголовок структурного элемента — заглавными и по центру; его подразделы
        тоже без номеров.
        """
        if number is None:
            text = " ".join(section.title.split()).upper() if level == 1 else section.title
            heading = document.add_paragraph(text, style=f"Heading {level}")
            if level == 1:
                _centered(heading)
        else:
            document.add_paragraph(f"{number}. {section.title}", style=f"Heading {level}")
        for text in section.paragraphs:
            document.add_paragraph(text)
        for sub_number, subsection in enumerate(section.subsections, start=1):
            sub = None if number is None else f"{number}.{sub_number}"
            self._add_section(document, subsection, sub, level=level + 1)

    def _add_references(self, document: DocxDocument, references: list[str]) -> None:
        """«СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ» — последним, без номера, пункты нумеруются."""
        _centered(document.add_paragraph(REFERENCES_TITLE, style="Heading 1"))
        for number, reference in enumerate(references, start=1):
            document.add_paragraph(f"{number}. {reference}")


# --- OOXML внутри python-docx -----------------------------------------------------------


def _centered(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)


def _add_field(paragraph: Paragraph, instruction: str, *, dirty: bool = False) -> None:
    """Поле Word (begin / instrText / separate / end) без заранее вычисленного значения."""

    def char(kind: str) -> BaseOxmlElement:
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), kind)
        if kind == "begin" and dirty:
            element.set(qn("w:dirty"), "true")
        return element

    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = f" {instruction} "
    for element in (char("begin"), code, char("separate"), char("end")):
        paragraph.add_run()._r.append(element)


def _update_fields_on_open(document: DocxDocument) -> None:
    """``w:updateFields`` на своём месте последовательности CT_Settings схемы OOXML."""
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        following = {qn(name) for name in _AFTER_UPDATE_FIELDS}
        anchor = next((child for child in settings if child.tag in following), None)
        if anchor is None:
            settings.append(update)
        else:
            anchor.addprevious(update)
    update.set(qn("w:val"), "true")


# Элементы w:settings, которые по схеме идут после w:updateFields (расширения w14/w15
# в конце последовательности — тоже после).
_AFTER_UPDATE_FIELDS = (
    "w:hdrShapeDefaults",
    "w:footnotePr",
    "w:endnotePr",
    "w:compat",
    "w:docVars",
    "w:rsids",
    "m:mathPr",
    "w:attachedSchema",
    "w:themeFontLang",
    "w:clrSchemeMapping",
    "w:doNotIncludeSubdocsInStats",
    "w:doNotAutoCompressPictures",
    "w:forceUpgrade",
    "w:captions",
    "w:readModeInkLockDown",
    "w:smartTagType",
    "w:shapeDefaults",
    "w:doNotEmbedSmartTags",
    "w:decimalSymbol",
    "w:listSeparator",
    "w14:docId",
    "w14:defaultImageDpi",
)


def _set_language(run_properties: BaseOxmlElement, language: str) -> None:
    lang = run_properties.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run_properties.append(lang)
    lang.set(qn("w:val"), language)


def _set_metadata(document: DocxDocument, title: str) -> None:
    """Свойства файла: только название; без автора, программы и дат шаблона."""
    properties = document.core_properties
    properties.title = title
    properties.author = ""
    properties.comments = ""
    properties.last_modified_by = ""
    properties.revision = 1
    element = properties._element
    for name in ("dcterms:created", "dcterms:modified"):
        found = element.find(qn(name))
        if found is not None:
            element.remove(found)


# Части шаблона python-docx, которые не относятся к документу и вводят в заблуждение:
# миниатюра пустого шаблона, свойства приложения («Microsoft Macintosh Word», 0 слов),
# хранилище библиографии со статичным GUID, копия стилей для Word 2010.
_STYLES_WITH_EFFECTS = "http://schemas.microsoft.com/office/2007/relationships/stylesWithEffects"
_TEMPLATE_ARTIFACT_RELATIONSHIPS = frozenset(
    {RT.THUMBNAIL, RT.EXTENDED_PROPERTIES, RT.CUSTOM_XML, _STYLES_WITH_EFFECTS}
)
# Идентификатор документа шаблона: одинаковый у всех отчётов.
_TEMPLATE_SETTINGS_IDS = ("w14:docId",)


def _strip_template_artifacts(document: DocxDocument) -> None:
    """Убирает артефакты шаблона, снимая связи на них.

    Части в пакет попадают обходом графа связей, поэтому без связи часть не
    сохраняется, а [Content_Types].xml и .rels остаются согласованными. Нумерация
    списков, тема, таблица шрифтов, веб-настройки и rsid остаются: на них ссылаются
    стили шаблона.
    """
    for source in (document.part.package, document.part):
        artifacts = [
            r_id
            for r_id, relationship in source.rels.items()
            if relationship.reltype in _TEMPLATE_ARTIFACT_RELATIONSHIPS
        ]
        for r_id in artifacts:
            source.rels.pop(r_id)
    settings = document.settings.element
    for name in _TEMPLATE_SETTINGS_IDS:
        for element in settings.findall(qn(name)):
            settings.remove(element)


def _normalize_zip(data: bytes) -> bytes:
    """Те же части в том же порядке, но с фиксированным временем записей ZIP."""
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            info = zipfile.ZipInfo(item.filename, date_time=_ZIP_TIMESTAMP)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            target.writestr(info, source.read(item.filename))
    return output.getvalue()
